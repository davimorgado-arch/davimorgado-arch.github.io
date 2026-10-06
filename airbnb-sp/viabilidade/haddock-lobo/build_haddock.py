"""Estudo de caso: SCP Haddock Lobo (studio ~18 m²) - fluxo, Airbnb, break-even e comparação com CDI.

Abas: Resumo (4 formas de pagamento lado a lado), Premissas, Tabela SCP (extraída do PDF),
Fluxo mensal (out/26 a dez/55, tudo em fórmulas).
"""
import json
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference

OUT = "Haddock_Lobo_estudo_de_caso.xlsx"
U = json.load(open("unidades.json"))

F = "Arial"
f_norm = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True); f_input = Font(name=F, size=10, color="0000FF")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF"); f_note = Font(name=F, size=9, color="6B6A64")
fill_key = PatternFill("solid", fgColor="FFFF00"); fill_head = PatternFill("solid", fgColor="1F2A2E")
fill_sub = PatternFill("solid", fgColor="EDEAE3"); fill_hl = PatternFill("solid", fgColor="F3DFD5")
thin = Side(style="thin", color="D0CCC2")
BRL = 'R$ #,##0;(R$ #,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; DT = 'mmm/yy'

wb = Workbook()

# ---------------- Tabela SCP ----------------
wt = wb.active; wt.title = "Tabela SCP"
wt["A1"] = "SCP Haddock Lobo - tabela out/26 (extraída do PDF enviado pelo usuário)"; wt["A1"].font = f_title
hdr = ["Unidade", "Tipologia", "Andar", "Área (m²)", "Previsão lançamento mar/27", "SCP à vista out/26",
       "3x SCP", "Ato 20% + 43x SCP", "Promo à vista", "SCP R$/m² à vista", "Desconto SCP vs lançamento"]
for j, h in enumerate(hdr, 1):
    c = wt.cell(row=3, column=j, value=h); c.font = f_bold; c.fill = fill_sub
    c.alignment = Alignment(wrap_text=True, vertical="top")
for i, u in enumerate(U):
    r = 4 + i
    vals = [u["unidade"], u["tipologia"], u["andar"], u["area"], u["lanc"], u["avista"], u["x3"], u["x43"], u["promo"], u["m2"]]
    for j, v in enumerate(vals, 1):
        c = wt.cell(row=r, column=j, value=v); c.font = f_norm
        if j >= 5: c.number_format = BRL
        if j == 4: c.number_format = '0.00'
    wt.cell(row=r, column=11, value=f"=1-MIN(F{r},IF(I{r}>0,I{r},F{r}))/E{r}").number_format = PCT
    wt.cell(row=r, column=11).font = f_norm
TL = 3 + len(U)
wt.cell(row=TL + 2, column=1, value="Fonte: Tabela_SCP_Haddock_Lobo_Out26.pdf (enviada pelo usuário). Valores copiados sem alteração; coluna K calculada.").font = Font(name=F, size=9, italic=True)
for j, w in enumerate([9, 15, 7, 9, 15, 13, 12, 14, 12, 12, 12], 1):
    wt.column_dimensions[get_column_letter(j)].width = w
wt.row_dimensions[3].height = 42; wt.freeze_panes = "B4"
T = "'Tabela SCP'"
def tcol(c): return f"{T}!${c}$4:${c}${TL}"

# ---------------- Premissas ----------------
wp = wb.create_sheet("Premissas", 0)
wp["A1"] = "Premissas - SCP Haddock Lobo"; wp["A1"].font = f_title
wp["A2"] = "Edite só as células em azul (amarelo = premissas que mais mexem no resultado)."; wp["A2"].font = Font(name=F, size=9, italic=True)
P = {}; row = 4
def section(t):
    global row
    row += 1; wp[f"A{row}"] = t; wp[f"A{row}"].font = f_bold
    for col in "ABCD": wp[f"{col}{row}"].fill = fill_sub
    row += 1
def inp(k, lab, v, fmt=None, note="", key=False, formula=False):
    global row
    wp[f"A{row}"] = lab; wp[f"A{row}"].font = f_norm
    c = wp[f"C{row}"]; c.value = v; c.font = f_norm if formula else f_input
    if key and not formula: c.fill = fill_key
    if fmt: c.number_format = fmt
    if note: wp[f"D{row}"] = note; wp[f"D{row}"].font = f_note
    P[k] = f"Premissas!$C${row}"; row += 1

section("1. Unidade")
inp("un", "Unidade escolhida (lista)", "308", note="Padrão: 308 (18,09 m², tem promo de R$ 250 mil = menor R$/m² entre os studios de ~18 m²)", key=True)
def look(col): return f"=INDEX({tcol(col)},MATCH({P['un']},{tcol('A')},0))"
inp("area", "Área (m²)", look("D"), '0.00', formula=True)
inp("lanc", "Previsão de lançamento (mar/27)", look("E"), BRL, formula=True)
inp("avista", "SCP à vista (out/26)", look("F"), BRL, formula=True)
inp("x3", "3x SCP (total)", look("G"), BRL, formula=True)
inp("x43", "Ato 20% + 43x SCP (total)", look("H"), BRL, formula=True)
inp("promo_raw", "Promo à vista (0 = sem promo para esta unidade)", f"=N({look('I')[1:]})", BRL, formula=True)
inp("promo", "Promo à vista usada no estudo", f"=IF({P['promo_raw']}>0,{P['promo_raw']},{P['avista']})", BRL, "Se a unidade não tem promo, usa o preço SCP à vista", formula=True)

section("2. Datas")
inp("compra", "Data da compra (adesão à SCP)", date(2026, 10, 1), DT, "Tabela válida em out/26", key=True)
inp("lanc_dt", "Lançamento previsto", date(2027, 3, 1), DT, "Tabela: previsão mar/27")
inp("entrega", "Entrega das chaves", date(2030, 6, 1), DT, "PREMISSA: a tabela não informa. Lançamento mar/27 + ~39 meses; o 43x termina em mai/30", key=True)
inp("mob_m", "Meses para mobiliar", 2, '0')
inp("ini_op", "Início do Airbnb", f"=DATE(YEAR({P['entrega']}),MONTH({P['entrega']})+{P['mob_m']}+1,1)", DT, formula=True)

section("3. Correções e custos de aquisição")
inp("corr", "Correção das parcelas SCP (a.a.)", 0.055, PCT, "PREMISSA: a tabela não informa. Usei INCC ~5,5%. Se for valor fixo, use 0", key=True)
inp("incc", "INCC até a entrega (corrige o preço de lançamento)", 0.055, PCT)
inp("vm_fator", "Valor de mercado na entrega (% da tabela de lançamento corrigida)", 0.80, PCT,
    "PREMISSA-CHAVE: revenda de studio costuma sair abaixo da tabela da incorporadora. 100% = vale o preço de lançamento", key=True)
inp("itbi", "ITBI + registro na entrega (% do valor de mercado)", 0.04, PCT, "Entrega da unidade ao sócio tende a ser dação em pagamento: ITBI sobre o valor de referência")
inp("mobilia", "Mobília, decoração e enxoval", 25000, BRL, "Studio de ~18 m²", key=True)

section("4. Operação Airbnb (valores de hoje)")
inp("diaria", "Diária média", 250, BRL, "Jardins/Consolação, studio: R$ 220-320 (estudo). Haddock Lobo fica entre Paulista e Oscar Freire", key=True)
inp("ocup", "Ocupação média", 0.62, PCT, "Mediana SP 61-63%", key=True)
inp("dias", "Dias por mês", 30, '0')
inp("ramp_m", "Meses de ramp-up", 3, '0'); inp("ramp_f", "Ocupação no ramp-up (% da média)", 0.6, PCT)
inp("tx_airbnb", "Taxa Airbnb", 0.16, PCT); inp("gestora", "Gestora (% da receita)", 0.20, PCT)
inp("cond", "Condomínio + IPTU (mês)", 600, BRL, "Estimativa para ~18 m² em prédio novo com amenities", key=True)
inp("contas", "Contas (luz, internet, streaming, amenities)", 280, BRL)
inp("manut", "Manutenção e reposição", 150, BRL)
inp("infl", "Inflação de diária e custos (a.a.)", 0.04, PCT)
inp("ir", "IR sobre aluguel (alíquota marginal)", 0.275, PCT, "Carnê-leão; deduz Airbnb, gestora, condomínio e IPTU")

section("5. Valorização e venda")
inp("valoriz", "Valorização após a entrega (a.a.)", 0.05, PCT, "SP: +3,8% em 12 meses (FipeZap jul/26)", key=True)
inp("corret", "Corretagem na venda", 0.06, PCT); inp("ir_gc", "IR sobre ganho de capital", 0.15, PCT)

section("6. Alternativa: aplicar o mesmo dinheiro (CDI, carteira espelho)")
inp("cdi26", "CDI médio 2026 (a.a.)", 0.14, PCT, "Selic ~14% em out/26")
inp("cdi27", "CDI médio 2027", 0.125, PCT, "Focus: Selic 12% fim/27")
inp("cdi28", "CDI médio 2028", 0.11, PCT, "Focus: 10,5% fim/28")
inp("cdi29", "CDI médio 2029 em diante", 0.10, PCT, "Focus: 10% em 2029", key=True)
inp("ir_ap", "IR sobre rendimento", 0.15, PCT)

section("Taxas mensais (calculadas)")
for k, lab in [("corr", "Correção SCP"), ("incc", "INCC"), ("infl", "Inflação"), ("valoriz", "Valorização")]:
    inp(k + "_m", f"{lab} ao mês", f"=(1+{P[k]})^(1/12)-1", '0.000%', formula=True)
inp("vm_entrega", "Valor de mercado estimado na entrega",
    f"={P['lanc']}*(1+{P['incc_m']})^((YEAR({P['entrega']})-YEAR({P['lanc_dt']}))*12+MONTH({P['entrega']})-MONTH({P['lanc_dt']}))*{P['vm_fator']}", BRL, formula=True)
inp("vm_m2", "… por m² (compare: Jardim Paulista ~R$ 17,8 mil/m² em 2026)", f"={P['vm_entrega']}/{P['area']}", BRL, formula=True)

wp.column_dimensions["A"].width = 52; wp.column_dimensions["B"].width = 2
wp.column_dimensions["C"].width = 16; wp.column_dimensions["D"].width = 92
dv = DataValidation(type="list", formula1=f"={tcol('A')}"); wp.add_data_validation(dv)
dv.add(P["un"].split("!")[1].replace("$", ""))

# ---------------- Fluxo mensal ----------------
wf = wb.create_sheet("Fluxo mensal")
wf["A1"] = "Fluxo mensal - SCP Haddock Lobo (R$ nominais; saídas negativas)"; wf["A1"].font = f_title
OPT = [("Promo à vista", "promo"), ("À vista", "avista"), ("3x", "x3"), ("20% + 43x", "x43")]
common = [("A","Mês",9),("B","Fase",11),("C","Meses desde a compra",8),("D","Ocupação",8),("E","Receita bruta",11),
          ("F","Taxa Airbnb",10),("G","Gestora",10),("H","Condomínio + IPTU",10),("I","Contas",9),("J","Manutenção",9),
          ("K","IR",9),("L","Resultado do imóvel (após IR)",12),("M","ITBI + registro",11),("N","Mobília",10),
          ("O","Valor de mercado",13),("P","CDI líquido (mês)",8)]
HR, first = 4, 5
months = []; d = date(2026, 10, 1)
while d <= date(2055, 12, 1):
    months.append(d); d = date(d.year + d.month // 12, d.month % 12 + 1, 1)
last = first + len(months) - 1
blocks = {}
col_i = 17  # Q
for name, key in OPT:
    cols = {}
    for suf in ("Pagamento SCP", "Fluxo de caixa", "Caixa acumulado", "Patrimônio se vender", "Posição total", "Saldo em CDI (espelho)"):
        cols[suf] = get_column_letter(col_i); col_i += 1
    blocks[key] = (name, cols)
heads = list(common)
for key, (name, cols) in blocks.items():
    for suf, L in cols.items(): heads.append((L, f"{name}: {suf}", 12))
for col, h, w in heads:
    c = wf[f"{col}{HR}"]; c.value = h; c.font = f_head; c.fill = fill_head
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center"); wf.column_dimensions[col].width = w
wf.row_dimensions[HR].height = 66
p = P
for idx, m in enumerate(months):
    r = first + idx; A = f"$A{r}"
    wf[f"A{r}"] = f"={p['compra']}" if idx == 0 else f"=DATE(YEAR(A{r-1}),MONTH(A{r-1})+1,1)"
    wf[f"B{r}"] = f'=IF({A}<{p["entrega"]},"Obra",IF({A}<{p["ini_op"]},"Entrega/mobília","Airbnb"))'
    wf[f"C{r}"] = f"=(YEAR({A})-YEAR({p['compra']}))*12+MONTH({A})-MONTH({p['compra']})"
    mop = f"((YEAR({A})-YEAR({p['ini_op']}))*12+MONTH({A})-MONTH({p['ini_op']}))"
    inf = f"(1+{p['infl_m']})^$C{r}"
    wf[f"D{r}"] = f"=IF({A}>={p['ini_op']},{p['ocup']}*IF({mop}<{p['ramp_m']},{p['ramp_f']},1),0)"
    wf[f"E{r}"] = f"=D{r}*{p['diaria']}*{p['dias']}*{inf}"
    wf[f"F{r}"] = f"=E{r}*{p['tx_airbnb']}"; wf[f"G{r}"] = f"=E{r}*{p['gestora']}"
    wf[f"H{r}"] = f"=IF({A}>={p['entrega']},{p['cond']}*{inf},0)"
    wf[f"I{r}"] = f"=IF({A}>={p['ini_op']},{p['contas']}*{inf},0)"
    wf[f"J{r}"] = f"=IF({A}>={p['ini_op']},{p['manut']}*{inf},0)"
    wf[f"K{r}"] = f"=IF({A}>={p['ini_op']},MAX(0,E{r}-F{r}-G{r}-H{r})*{p['ir']},0)"
    wf[f"L{r}"] = f"=E{r}-F{r}-G{r}-H{r}-I{r}-J{r}-K{r}"
    wf[f"M{r}"] = f"=({A}={p['entrega']})*{p['itbi']}*{p['vm_entrega']}"
    wf[f"N{r}"] = f"=IF(AND({A}>={p['entrega']},{A}<{p['ini_op']}),{p['mobilia']}/MAX(1,{p['mob_m']}),0)"
    me = f"((YEAR({A})-YEAR({p['entrega']}))*12+MONTH({A})-MONTH({p['entrega']}))"
    wf[f"O{r}"] = f'=IF({A}>={p["entrega"]},{p["vm_entrega"]}*(1+{p["valoriz_m"]})^{me},"")'
    cdi = f"IF(YEAR({A})<=2026,{p['cdi26']},IF(YEAR({A})=2027,{p['cdi27']},IF(YEAR({A})=2028,{p['cdi28']},{p['cdi29']})))"
    wf[f"P{r}"] = f"=(1+{cdi}*(1-{p['ir_ap']}))^(1/12)-1"
    corr = f"(1+{p['corr_m']})^$C{r}"
    pay = {
        "promo": f"=({A}={p['compra']})*{p['promo']}",
        "avista": f"=({A}={p['compra']})*{p['avista']}",
        "x3": f"=AND({A}>={p['compra']},{A}<EDATE({p['compra']},3))*{p['x3']}/3*{corr}",
        "x43": f"=({A}={p['compra']})*0.2*{p['x43']}+AND({A}>{p['compra']},{A}<=EDATE({p['compra']},43))*0.8*{p['x43']}/43*{corr}",
    }
    for key, (name, cols) in blocks.items():
        pg, fl, ac, pt, po, cd = cols.values()
        wf[f"{pg}{r}"] = pay[key]
        wf[f"{fl}{r}"] = f"=L{r}-M{r}-N{r}-{pg}{r}"
        wf[f"{ac}{r}"] = f"={fl}{r}" if idx == 0 else f"={ac}{r-1}+{fl}{r}"
        custo = f"(SUM(${pg}${first}:${pg}${last})+{p['itbi']}*{p['vm_entrega']})"
        wf[f"{pt}{r}"] = (f'=IF({A}>={p["entrega"]},O{r}*(1-{p["corret"]})'
                          f'-{p["ir_gc"]}*MAX(0,O{r}*(1-{p["corret"]})-{custo}),"")')
        wf[f"{po}{r}"] = f'=IF({A}>={p["entrega"]},{ac}{r}+{pt}{r},"")'
        prev = f"{cd}{r-1}" if idx > 0 else "0"
        wf[f"{cd}{r}"] = f"={prev}*(1+$P{r})-{fl}{r}"
    for col, _, _ in heads:
        c = wf[f"{col}{r}"]; c.font = f_norm
        c.number_format = DT if col == "A" else ('0' if col == "C" else (PCT if col == "D" else ('0.000%' if col == "P" else BRL)))
wf.freeze_panes = f"C{first}"
FL = "'Fluxo mensal'"
def R(col): return f"{FL}!${col}${first}:${col}${last}"

# ---------------- Resumo ----------------
ws = wb.create_sheet("Resumo", 0)
ws["A1"] = "SCP Haddock Lobo - estudo de caso do studio"; ws["A1"].font = f_title
ws["A2"] = ("Compara as 4 formas de pagamento da tabela. Sem financiamento: tudo é pago até a entrega. "
            "Imóvel vendido (líquido) x mesmo dinheiro aplicado no CDI (líquido de IR).")
ws["A2"].font = Font(name=F, size=9, italic=True)
info = [("Unidade", f"={P['un']}", '@'), ("Área (m²)", f"={P['area']}", '0.00'),
        ("Previsão de lançamento (mar/27)", f"={P['lanc']}", BRL),
        ("Valor de mercado estimado na entrega", f"={P['vm_entrega']}", BRL),
        ("… R$/m² na entrega", f"={P['vm_m2']}", BRL),
        ("Entrega das chaves (premissa)", f"={P['entrega']}", DT),
        ("Início do Airbnb", f"={P['ini_op']}", DT),
        ("Resultado mensal do imóvel após IR (1º mês pós ramp-up)", f"=INDEX({R('L')},MATCH(EDATE({P['ini_op']},{P['ramp_m']}),{R('A')},0))", BRL)]
r = 4
for lab, fml, fmt in info:
    ws[f"A{r}"] = lab; ws[f"A{r}"].font = f_norm
    ws[f"B{r}"] = fml; ws[f"B{r}"].number_format = fmt; ws[f"B{r}"].font = f_bold; r += 1
r += 1
ws.cell(row=r, column=1, value="Indicador").font = f_head; ws.cell(row=r, column=1).fill = fill_head
for j, (name, key) in enumerate(OPT):
    c = ws.cell(row=r, column=2 + j, value=name); c.font = f_head; c.fill = fill_head
    c.alignment = Alignment(horizontal="center")
r += 1
NA = '"Não atinge até 2055"'
def at(col, months_after):
    return f"INDEX({R(col)},MATCH(EDATE({P['entrega']},{months_after}),{R('A')},0))"
rows_r = [
 ("Total pago à SCP", lambda c: f"=SUM({R(c['Pagamento SCP'])})", BRL),
 ("Desembolso total até o início do Airbnb (SCP + ITBI + mobília)", lambda c: f"=-INDEX({R(c['Caixa acumulado'])},MATCH({P['ini_op']},{R('A')},0))", BRL),
 ("Ganho no papel na entrega (posição se vender na entrega)", lambda c: f"={at(c['Posição total'],0)}", BRL),
 ("Retorno do aluguel sobre o desembolso (a.a., 1º ano pleno)", lambda c: f"=12*INDEX({R('L')},MATCH(EDATE({P['ini_op']},{P['ramp_m']}),{R('A')},0))/-INDEX({R(c['Caixa acumulado'])},MATCH({P['ini_op']},{R('A')},0))", PCT),
 ("Break-even de caixa (aluguel devolve tudo que saiu)", lambda c: f'=IFERROR(INDEX({R("A")},MATCH(1,INDEX(({R(c["Caixa acumulado"])}>=0)*({R("A")}>={P["ini_op"]}),0),0)),{NA})', DT),
 ("Imóvel: posição se vender 5 anos após a entrega", lambda c: f"={at(c['Posição total'],60)}", BRL),
 ("CDI: saldo líquido no mesmo momento (5 anos)", lambda c: f"={at(c['Saldo em CDI (espelho)'],60)}+INDEX({R(c['Caixa acumulado'])},MATCH(EDATE({P['entrega']},60),{R('A')},0))", BRL),
 ("Imóvel: posição se vender 10 anos após a entrega", lambda c: f"={at(c['Posição total'],120)}", BRL),
 ("CDI: saldo líquido no mesmo momento (10 anos)", lambda c: f"={at(c['Saldo em CDI (espelho)'],120)}+INDEX({R(c['Caixa acumulado'])},MATCH(EDATE({P['entrega']},120),{R('A')},0))", BRL),
 ("Imóvel: posição se vender 20 anos após a entrega", lambda c: f"={at(c['Posição total'],240)}", BRL),
 ("CDI: saldo líquido no mesmo momento (20 anos)", lambda c: f"={at(c['Saldo em CDI (espelho)'],240)}+INDEX({R(c['Caixa acumulado'])},MATCH(EDATE({P['entrega']},240),{R('A')},0))", BRL),
]
start_tab = r
for lab, fn, fmt in rows_r:
    ws.cell(row=r, column=1, value=lab).font = f_norm
    for j, (name, key) in enumerate(OPT):
        c = ws.cell(row=r, column=2 + j, value=fn(blocks[key][1])); c.number_format = fmt; c.font = f_norm
        c.alignment = Alignment(horizontal="right")
    if lab.startswith("CDI"):
        for j in range(5): ws.cell(row=r, column=1 + j).fill = fill_hl
    r += 1
r += 1
notes = [
 "Como ler:",
 "• Posição total = caixa acumulado (tudo que saiu e entrou) + o que sobraria vendendo (após corretagem e IR sobre ganho).",
 "• Linhas CDI: o mesmo dinheiro tirado do bolso aplicado no CDI (líquido de 15% de IR), na mesma régua da posição do imóvel.",
 "• Riscos fora da planilha: SCP antes do registro de incorporação (sem patrimônio de afetação), aprovação do projeto, atraso, regulação CVM, liquidez da cota.",
 "• Datas e correção das parcelas não constam da tabela: confirmar com a incorporadora (abas Premissas, seções 2 e 3).",
]
for t in notes:
    ws.cell(row=r, column=1, value=t).font = Font(name=F, size=9, bold=t.endswith(":")); r += 1
ws.column_dimensions["A"].width = 60
for col in "BCDE": ws.column_dimensions[col].width = 16

# gráfico: posição do imóvel (promo) x CDI espelho
ch = LineChart(); ch.title = "Promo à vista: imóvel (posição) x CDI (saldo + caixa)"; ch.height = 9; ch.width = 20
cp = blocks["promo"][1]
# série auxiliar do CDI na mesma régua (saldo + caixa acumulado) na coluna à direita do fluxo
aux_col = get_column_letter(col_i)
wf[f"{aux_col}{HR}"] = "Promo: CDI na régua da posição (saldo + caixa)"; wf[f"{aux_col}{HR}"].font = f_head; wf[f"{aux_col}{HR}"].fill = fill_head
wf[f"{aux_col}{HR}"].alignment = Alignment(wrap_text=True, horizontal="center")
wf.column_dimensions[aux_col].width = 14
for idx in range(len(months)):
    rr = first + idx
    wf[f"{aux_col}{rr}"] = f"={cp['Saldo em CDI (espelho)']}{rr}+{cp['Caixa acumulado']}{rr}"
    wf[f"{aux_col}{rr}"].number_format = BRL; wf[f"{aux_col}{rr}"].font = f_norm
from openpyxl.utils import column_index_from_string as ci
ch.add_data(Reference(wf, min_col=ci(cp["Posição total"]), min_row=HR, max_row=last), titles_from_data=True)
ch.add_data(Reference(wf, min_col=ci(aux_col), min_row=HR, max_row=last), titles_from_data=True)
ch.set_categories(Reference(wf, min_col=1, min_row=first, max_row=last))
ch.y_axis.number_format = 'R$ #,##0'; ch.x_axis.number_format = 'yyyy'; ch.x_axis.tickLblSkip = 24
ch.series[0].graphicalProperties.line.solidFill = "B4441F"; ch.series[1].graphicalProperties.line.solidFill = "2F6B3A"
ws.add_chart(ch, "G4")

wb.save(OUT); print("ok", OUT)
